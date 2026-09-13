"""Сборка сопроводительных документов к защите: markdown → DOCX (python-docx).

Тексты лежат в ``slides/content/*.md``, этот скрипт только рендерит:

    .venv/Scripts/python slides/build_docs.py

Выход (в ``slides/``): «Речь доклада», «Термины и определения», «Вопросы и ответы»,
«Сценарий демонстрации» — .docx, A4, с нумерацией страниц.

Поддерживаемая разметка: ``#``/``##``/``###`` заголовки, ``> цитата``, абзацы,
маркированные и нумерованные списки, markdown-таблицы, блоки кода ``` и
инлайновые ``**жирный**``, ``*курсив*``, ```код```.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
CONTENT = HERE / "content"

# Документы: файл контента → имя выходного файла
DOCS = {
    "speech.md": "Речь доклада.docx",
    "terms.md": "Термины и определения.docx",
    "qa.md": "Вопросы и ответы.docx",
    "demo-script.md": "Сценарий демонстрации.docx",
}

FONT = "Segoe UI"
FONT_MONO = "Consolas"
INK = RGBColor(0x1B, 0x27, 0x38)
ACCENT = RGBColor(0x1B, 0x4F, 0xA8)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
CODE_INK = RGBColor(0xA8, 0x3A, 0x0C)
SHADE = "F2F5FA"

RE_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
RE_BULLET = re.compile(r"^(?P<indent>\s*)[-*]\s+(?P<text>.+)$")
RE_NUMBER = re.compile(r"^\s*\d+\.\s+(?P<text>.+)$")


# --- низкоуровневые помощники -------------------------------------------------


def _shade(element, fill: str) -> None:
    """Заливка фона абзаца или ячейки таблицы."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _add_runs(paragraph, text: str, *, size: float, color: RGBColor = INK, bold: bool = False) -> None:
    """Разложить инлайновую разметку на отдельные run-ы."""
    for part in RE_INLINE.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        if part.startswith("**") and part.endswith("**"):
            inner = part[2:-2]
            run.text = inner.replace("`", "")
            run.font.bold = True
            if "`" in inner:
                run.font.name = FONT_MONO
                run.font.size = Pt(size - 0.5)
                run.font.color.rgb = CODE_INK
        elif part.startswith("*") and part.endswith("*"):
            run.text = part[1:-1]
            run.font.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = FONT_MONO
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = CODE_INK
        else:
            run.text = part


def _page_numbers(document: Document) -> None:
    """Нумерация страниц в колонтитуле («стр. N»)."""
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("стр. ")
    run.font.name = FONT
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    field = paragraph.add_run()
    field.font.name = FONT
    field.font.size = Pt(9)
    field.font.color.rgb = MUTED
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field._r.append(begin)
    field._r.append(instr)
    field._r.append(end)


def _setup(document: Document) -> None:
    """Страница A4, поля, базовый шрифт, стили заголовков."""
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before in (
        ("Title", 21, INK, 0),
        ("Heading 1", 15, ACCENT, 16),
        ("Heading 2", 12.5, INK, 12),
    ):
        style = document.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    _page_numbers(document)


# --- блоки --------------------------------------------------------------------


def _add_quote(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(4)
    _shade(paragraph._p.get_or_add_pPr(), SHADE)
    _add_runs(paragraph, " ".join(lines), size=10.5, color=MUTED)
    for run in paragraph.runs:
        run.font.italic = True


def _add_code(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05
    _shade(paragraph._p.get_or_add_pPr(), SHADE)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line or " ")
        run.font.name = FONT_MONO
        run.font.size = Pt(9.5)
        run.font.color.rgb = INK


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            _add_runs(paragraph, text, size=10, bold=(r == 0))
            if r == 0:
                _shade(cell._tc.get_or_add_tcPr(), SHADE)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _parse_table(rows: list[str]) -> list[list[str]]:
    cells = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
    if len(cells) > 1 and all(set(c) <= set("-: ") and c for c in cells[1]):
        del cells[1]
    return cells


def render(source: Path, target: Path) -> Path:
    """Собрать один DOCX из markdown-файла."""
    document = Document()
    _setup(document)

    lines = source.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)
    paragraph_buffer: list[str] = []

    def flush() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(paragraph, " ".join(paragraph_buffer), size=11.5)
            paragraph_buffer = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush()
            i += 1
            continue

        if stripped.startswith("# "):
            flush()
            document.add_paragraph(stripped[2:], style="Title")
            i += 1
            continue

        if stripped.startswith("## "):
            flush()
            document.add_paragraph(stripped[3:], style="Heading 1")
            i += 1
            continue

        if stripped.startswith("### "):
            flush()
            document.add_paragraph(stripped[4:], style="Heading 2")
            i += 1
            continue

        if stripped.startswith("> "):
            flush()
            quote: list[str] = []
            while i < n and lines[i].strip().startswith("> "):
                quote.append(lines[i].strip()[2:])
                i += 1
            _add_quote(document, quote)
            continue

        if stripped.startswith("```"):
            flush()
            code: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            _add_code(document, code)
            continue

        if stripped.startswith("|"):
            flush()
            table: list[str] = []
            while i < n and lines[i].strip().startswith("|"):
                table.append(lines[i].strip())
                i += 1
            _add_table(document, _parse_table(table))
            continue

        bullet = RE_BULLET.match(line)
        number = RE_NUMBER.match(line)
        if bullet or number:
            flush()
            while i < n:
                bullet = RE_BULLET.match(lines[i])
                number = RE_NUMBER.match(lines[i])
                if not (bullet or number):
                    break
                if bullet:
                    level = len(bullet.group("indent")) // 2
                    style = "List Bullet" if level == 0 else "List Bullet 2"
                    text = bullet.group("text")
                else:
                    style = "List Number"
                    text = number.group("text")
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.space_after = Pt(3)
                _add_runs(paragraph, text, size=11.5)
                i += 1
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush()
    document.save(str(target))
    return target


def main() -> None:
    for name, output in DOCS.items():
        path = render(CONTENT / name, HERE / output)
        print(f"готово: {path.name}")


if __name__ == "__main__":
    main()
