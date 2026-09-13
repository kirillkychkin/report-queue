"""Мини-движок слайдов: markdown-контент -> PPTX (python-pptx).

Тексты слайдов лежат в ``slides/content/*.md``, скрипты ``build_talk.py`` /
``build_project.py`` только вызывают отсюда ``build_deck()``. Формат разметки
намеренно маленький и полностью описан в ``parse_deck()``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --- Тема оформления -------------------------------------------------------

INK = RGBColor(0x1B, 0x27, 0x38)  # основной текст
MUTED = RGBColor(0x5B, 0x6B, 0x7C)  # подписи, подпункты
ACCENT = RGBColor(0x2F, 0x6F, 0xED)  # акцент (синий)
ACCENT2 = RGBColor(0xE8, 0x59, 0x0C)  # второй акцент (оранжевый)
BG = RGBColor(0xFF, 0xFF, 0xFF)
BG_SOFT = RGBColor(0xF2, 0xF5, 0xFA)  # плашки, шапка таблицы
LINE = RGBColor(0xD5, 0xDD, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Segoe UI"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_TOP = Inches(1.62)
CONTENT_BOTTOM = Inches(6.92)
CONTENT_W = SLIDE_W - 2 * MARGIN


# --- Модель контента -------------------------------------------------------


@dataclass
class Block:
    """Один блок содержимого слайда."""

    kind: str  # bullets | table | image | code | callout
    data: Any


@dataclass
class Slide:
    title: str = ""
    subtitle: str = ""
    kind: str = "content"  # content | section
    blocks: list[Block] = field(default_factory=list)
    notes: str = ""


@dataclass
class Deck:
    meta: dict[str, str]
    slides: list[Slide]
    base_dir: Path


# --- Разбор markdown -------------------------------------------------------

RE_IMAGE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
RE_BULLET = re.compile(r"^(?P<indent>\s*)[-*]\s+(?P<text>.+?)\s*$")


def parse_deck(path: Path) -> Deck:
    """Разобрать файл с текстами слайдов.

    Формат:

    * front matter между ``---`` в начале файла — ``ключ: значение`` (титул, автор…);
    * ``## Заголовок`` начинает слайд;
    * ``> текст`` сразу после заголовка — подзаголовок;
    * ``::: section`` — слайд-разделитель;
    * ``- пункт`` / ``  - подпункт`` — маркированный список;
    * markdown-таблица (строки с ``|``) — таблица;
    * ``![alt](путь)`` — картинка (путь относительно файла контента);
    * ``` — блок кода;
    * ``!> текст`` — выделенная плашка с главной мыслью;
    * ``::: notes`` — всё до конца слайда уходит в заметки докладчика.
    """
    raw = path.read_text(encoding="utf-8")
    lines = raw.splitlines()
    meta: dict[str, str] = {}

    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        for line in lines[1:end]:
            if ":" in line:
                key, value = line.split(":", 1)
                meta[key.strip()] = value.strip()
        lines = lines[end + 1 :]

    slides: list[Slide] = []
    current: list[str] = []
    title = ""
    for line in lines:
        if line.startswith("## "):
            if title:
                slides.append(_parse_slide(title, current))
            title = line[3:].strip()
            current = []
        elif title:
            current.append(line)
    if title:
        slides.append(_parse_slide(title, current))

    return Deck(meta=meta, slides=slides, base_dir=path.parent)


def _parse_slide(title: str, body: list[str]) -> Slide:
    slide = Slide(title=title)
    i = 0
    n = len(body)
    while i < n:
        line = body[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "::: section":
            slide.kind = "section"
            i += 1
            continue

        if stripped == "::: notes":
            slide.notes = "\n".join(body[i + 1 :]).strip()
            break

        if stripped.startswith("> ") and not slide.blocks:
            slide.subtitle = stripped[2:].strip()
            i += 1
            continue

        if stripped.startswith("!> "):
            slide.blocks.append(Block("callout", stripped[3:].strip()))
            i += 1
            continue

        image = RE_IMAGE.match(stripped)
        if image:
            slide.blocks.append(Block("image", image.group("src")))
            i += 1
            continue

        if stripped.startswith("```"):
            code: list[str] = []
            i += 1
            while i < n and not body[i].strip().startswith("```"):
                code.append(body[i])
                i += 1
            i += 1
            slide.blocks.append(Block("code", "\n".join(code)))
            continue

        if stripped.startswith("|"):
            rows: list[str] = []
            while i < n and body[i].strip().startswith("|"):
                rows.append(body[i].strip())
                i += 1
            slide.blocks.append(Block("table", _parse_table(rows)))
            continue

        bullet = RE_BULLET.match(line)
        if bullet:
            items: list[tuple[int, str]] = []
            while i < n:
                match = RE_BULLET.match(body[i])
                if not match:
                    if body[i].strip():
                        break
                    i += 1
                    continue
                level = len(match.group("indent")) // 2
                items.append((min(level, 2), match.group("text")))
                i += 1
            slide.blocks.append(Block("bullets", items))
            continue

        # обычный абзац — как пункт без маркера
        slide.blocks.append(Block("bullets", [(-1, stripped)]))
        i += 1

    return slide


def _parse_table(rows: list[str]) -> list[list[str]]:
    cells = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
    # вторая строка markdown-таблицы — разделитель ---|---
    if len(cells) > 1 and all(set(c) <= set("-: ") and c for c in cells[1]):
        del cells[1]
    return cells


# --- Инлайновое форматирование --------------------------------------------

RE_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _add_runs(paragraph, text: str, size: Pt, color: RGBColor, bold: bool = False) -> None:
    """Разложить ``**жирный**`` и ```код``` на отдельные run-ы."""
    for part in RE_INLINE.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        run.font.size = size
        run.font.color.rgb = color
        run.font.name = FONT
        run.font.bold = bold
        if part.startswith("**") and part.endswith("**"):
            inner = part[2:-2]
            run.text = inner.replace("`", "")
            run.font.bold = True
            run.font.color.rgb = INK if color is MUTED else color
            if "`" in inner:  # жирный код: **`rate_limit`**
                run.font.name = FONT_MONO
                run.font.size = Pt(size.pt - 1)
                run.font.color.rgb = ACCENT2
        elif part.startswith("*") and part.endswith("*"):
            run.text = part[1:-1]
            run.font.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = FONT_MONO
            run.font.size = Pt(size.pt - 1)
            run.font.color.rgb = ACCENT2
        else:
            run.text = part


def _plain(text: str) -> str:
    return text.replace("**", "").replace("*", "").replace("`", "")


# --- Примитивы отрисовки ---------------------------------------------------


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return frame


def _rect(slide, left, top, width, height, fill: RGBColor | None, line: RGBColor | None = None):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def _set_bullet(paragraph, char: str, color: RGBColor, indent_in: float) -> None:
    """Проставить маркер списка через XML (python-pptx этого не умеет)."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(int(-Inches(0.24))))
    for tag, attrs in (
        ("a:buClr", None),
        ("a:buFont", {"typeface": "Arial"}),
        ("a:buChar", {"char": char}),
    ):
        element = pPr.makeelement(qn(tag), attrs or {})
        if tag == "a:buClr":
            element.append(pPr.makeelement(qn("a:srgbClr"), {"val": f"{color}"}))
        pPr.append(element)


# --- Отрисовка слайдов -----------------------------------------------------


def _blocks_height(blocks: list[Block], width: Emu, scale: float) -> Emu:
    """Суммарная высота всех блоков, кроме картинки (она занимает остаток)."""
    total = Emu(0)
    for block in blocks:
        if block.kind == "bullets":
            total += _bullets_height(block.data, width, scale) + Inches(0.12)
        elif block.kind == "table":
            total += _table_height(block.data, width, scale) + Inches(0.3)
        elif block.kind == "callout":
            total += _callout_height(block.data, width, scale) + Inches(0.18)
        elif block.kind == "code":
            total += _code_height(block.data, scale) + Inches(0.18)
    return total


def _fit_scale(blocks: list[Block], available: Emu, width: Emu) -> float:
    """Подобрать общий масштаб шрифтов так, чтобы содержимое влезло на слайд."""
    has_image = any(block.kind == "image" for block in blocks)
    budget = available - (Inches(2.4) if has_image else Inches(0))
    scale = 1.0
    while scale > 0.66 and _blocks_height(blocks, width, scale) > budget:
        scale -= 0.04
    return round(scale, 2)


def _vertical_offset(blocks, base_dir: Path, width: Emu, scale: float, available: Emu) -> Emu:
    """Немного опустить разреженный слайд, чтобы содержимое не липло к заголовку."""
    total = _blocks_height(blocks, width, scale)
    for block in blocks:
        if block.kind == "image":
            src = (base_dir / block.data).resolve()
            total += _image_height(src, width, available - total) + Inches(0.12)
    free = available - total
    return Emu(int(min(free / 2, Inches(0.9)))) if free > 0 else Emu(0)


def _new_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # пустой макет
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = BG
    return slide


def _draw_header(slide, title: str, subtitle: str) -> Emu:
    """Заголовок + акцентная черта. Возвращает верхнюю границу контента."""
    frame = _textbox(slide, MARGIN, Inches(0.48), CONTENT_W, Inches(0.85))
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    size = Pt(30) if len(_plain(title)) <= 52 else Pt(25)
    _add_runs(paragraph, title, size, INK, bold=True)

    _rect(slide, MARGIN, Inches(1.28), Inches(1.15), Pt(4), ACCENT)

    top = CONTENT_TOP
    if subtitle:
        sub = _textbox(slide, MARGIN, Inches(1.46), CONTENT_W, Inches(0.4))
        paragraph = sub.paragraphs[0]
        _add_runs(paragraph, subtitle, Pt(15), MUTED)
        top = Inches(2.0)
    return top


def _draw_footer(slide, footer: str, number: int) -> None:
    frame = _textbox(slide, MARGIN, Inches(7.02), CONTENT_W, Inches(0.3))
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = footer
    run.font.size = Pt(10)
    run.font.name = FONT
    run.font.color.rgb = MUTED

    frame = _textbox(slide, SLIDE_W - MARGIN - Inches(1.0), Inches(7.02), Inches(1.0), Inches(0.3))
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = str(number)
    run.font.size = Pt(10)
    run.font.name = FONT
    run.font.bold = True
    run.font.color.rgb = ACCENT


def _bullet_font_size(items: list[tuple[int, str]]) -> float:
    """Чем больше пунктов, тем мельче шрифт — чтобы всё влезло на слайд."""
    weight = sum(1 + len(_plain(text)) // 95 for _, text in items)
    if weight <= 5:
        return 20.0
    if weight <= 7:
        return 18.0
    if weight <= 9:
        return 16.5
    if weight <= 12:
        return 15.0
    return 13.5


def _wrapped_lines(text: str, size: float, width_in: float) -> int:
    """Сколько строк займёт текст в блоке заданной ширины."""
    chars_per_line = max(int(width_in * 72 / (size * 0.50)), 8)
    return max(1, -(-len(_plain(text)) // chars_per_line))


def _bullets_height(items, width: Emu, scale: float) -> Emu:
    size = _bullet_font_size(items) * scale
    total = 0.0
    for level, text in items:
        font = size if level <= 0 else size - 2.5 * scale
        lines = _wrapped_lines(text, font, width / Inches(1) - 0.3 * max(level, 0))
        total += lines * font * 1.32 / 72 + (0.10 if level <= 0 else 0.06)
    return Inches(total)


def _draw_bullets(slide, items, top: Emu, width: Emu, scale: float) -> Emu:
    size = _bullet_font_size(items) * scale
    height = _bullets_height(items, width, scale)
    frame = _textbox(slide, MARGIN, top, width, height)
    for index, (level, text) in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt((7 if level <= 0 else 4) * scale)
        if level == 0:
            _add_runs(paragraph, text, Pt(size), INK)
            _set_bullet(paragraph, "■", ACCENT, 0.26)
        elif level >= 1:
            _add_runs(paragraph, text, Pt(size - 2.5 * scale), MUTED)
            _set_bullet(paragraph, "–", ACCENT2, 0.62)
        else:  # абзац без маркера
            _add_runs(paragraph, text, Pt(size), INK)
    return top + height + Inches(0.12)


def _callout_height(text: str, width: Emu, scale: float) -> Emu:
    size = 16.0 * scale
    lines = _wrapped_lines(text, size, width / Inches(1) - 0.7)
    return Inches(0.30 + lines * size * 1.35 / 72)


def _draw_callout(slide, text: str, top: Emu, width: Emu, scale: float) -> Emu:
    height = _callout_height(text, width, scale)
    _rect(slide, MARGIN, top, width, height, BG_SOFT, LINE)
    _rect(slide, MARGIN, top, Pt(5), height, ACCENT2)
    frame = _textbox(slide, MARGIN + Inches(0.26), top + Inches(0.14), width - Inches(0.5), height)
    paragraph = frame.paragraphs[0]
    _add_runs(paragraph, text, Pt(16 * scale), INK)
    return top + height + Inches(0.18)


def _code_size(code: str, scale: float) -> float:
    lines = code.splitlines() or [""]
    base = 13.0 if max(len(line) for line in lines) <= 72 else 11.0
    return base * scale


def _code_height(code: str, scale: float) -> Emu:
    lines = code.splitlines() or [""]
    return Inches(0.26 + len(lines) * _code_size(code, scale) * 1.4 / 72)


def _draw_code(slide, code: str, top: Emu, width: Emu, scale: float) -> Emu:
    lines = code.splitlines() or [""]
    size = _code_size(code, scale)
    height = _code_height(code, scale)
    _rect(slide, MARGIN, top, width, height, BG_SOFT, LINE)
    frame = _textbox(slide, MARGIN + Inches(0.18), top + Inches(0.13), width - Inches(0.36), height)
    frame.word_wrap = False
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line or " "
        run.font.name = FONT_MONO
        run.font.size = Pt(size)
        run.font.color.rgb = INK
    return top + height + Inches(0.18)


def _table_layout(rows: list[list[str]], width: Emu, scale: float):
    """Ширины колонок и высоты строк с учётом переноса текста в ячейках."""
    n_cols = len(rows[0])
    size = (14.0 if len(rows) <= 6 else (12.5 if len(rows) <= 9 else 11.0)) * scale

    weights = [max(len(_plain(row[col])) for row in rows) + 6 for col in range(n_cols)]
    total_weight = sum(weights)
    col_widths = [Emu(int(width * weight / total_weight)) for weight in weights]

    row_heights = []
    for row in rows:
        lines = max(
            _wrapped_lines(text, size, col_widths[col] / Inches(1) - 0.18)
            for col, text in enumerate(row)
        )
        row_heights.append(Inches(lines * size * 1.32 / 72 + 0.19))
    return col_widths, row_heights, size


def _table_height(rows, width: Emu, scale: float) -> Emu:
    _, row_heights, _ = _table_layout(rows, width, scale)
    return Emu(int(sum(row_heights)))


def _draw_table(slide, rows: list[list[str]], top: Emu, width: Emu, scale: float) -> Emu:
    n_rows, n_cols = len(rows), len(rows[0])
    col_widths, row_heights, body_size = _table_layout(rows, width, scale)
    height = Emu(int(sum(row_heights)))
    shape = slide.shapes.add_table(n_rows, n_cols, MARGIN, top, width, height)
    table = shape.table
    table.first_row = True
    table.horz_banding = True

    for col in range(n_cols):
        table.columns[col].width = col_widths[col]

    for r, row in enumerate(rows):
        table.rows[r].height = row_heights[r]
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.09)
            cell.margin_right = Inches(0.09)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if r == 0 else (BG if r % 2 else BG_SOFT)
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            if r == 0:
                _add_runs(paragraph, text, Pt(body_size), WHITE, bold=True)
            else:
                _add_runs(paragraph, text, Pt(body_size), INK)
    return top + height + Inches(0.3)


def _png_size(src: Path) -> tuple[int, int]:
    """Размеры PNG из заголовка файла — без Pillow."""
    header = src.read_bytes()[:24]
    if header[:8] != bytes([137, 80, 78, 71, 13, 10, 26, 10]):
        return (16, 9)
    return (
        int.from_bytes(header[16:20], "big"),
        int.from_bytes(header[20:24], "big"),
    )


def _image_height(src: Path, width: Emu, available: Emu) -> Emu:
    px_w, px_h = _png_size(src)
    height = Emu(int(width * px_h / px_w))
    return min(height, available)


def _draw_image(slide, src: Path, top: Emu, width: Emu, bottom: Emu) -> Emu:
    picture = slide.shapes.add_picture(str(src), MARGIN, top, width=width)
    available = bottom - top
    if picture.height > available:
        scale = available / picture.height
        picture.height = Emu(int(picture.height * scale))
        picture.width = Emu(int(picture.width * scale))
    picture.left = Emu(int((SLIDE_W - picture.width) / 2))
    return top + picture.height + Inches(0.12)


def _draw_title_slide(prs: Presentation, meta: dict[str, str]) -> None:
    slide = _new_slide(prs)
    _rect(slide, Emu(0), Emu(0), SLIDE_W, Inches(2.95), RGBColor(0x12, 0x2B, 0x57))
    _rect(slide, Emu(0), Inches(2.95), SLIDE_W, Pt(6), ACCENT2)

    frame = _textbox(slide, MARGIN, Inches(0.85), CONTENT_W, Inches(1.6))
    paragraph = frame.paragraphs[0]
    _add_runs(paragraph, meta.get("title", ""), Pt(38), WHITE, bold=True)
    if meta.get("subtitle"):
        paragraph = frame.add_paragraph()
        paragraph.space_before = Pt(10)
        _add_runs(paragraph, meta["subtitle"], Pt(19), RGBColor(0xC3, 0xD4, 0xF5))

    rows = [
        ("Дисциплина", meta.get("course", "")),
        ("Докладчик", meta.get("author", "")),
        ("Репозиторий", meta.get("repo", "")),
        ("Дата", meta.get("date", "")),
    ]
    frame = _textbox(slide, MARGIN, Inches(3.75), CONTENT_W, Inches(3.0))
    first = True
    for label, value in rows:
        if not value:
            continue
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.space_after = Pt(12)
        run = paragraph.add_run()
        run.text = f"{label}:  "
        run.font.size = Pt(15)
        run.font.name = FONT
        run.font.color.rgb = MUTED
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(17)
        run.font.name = FONT if label != "Репозиторий" else FONT_MONO
        run.font.bold = label == "Докладчик"
        run.font.color.rgb = INK


def _draw_section_slide(slide, title: str, subtitle: str) -> None:
    _rect(slide, Emu(0), Emu(0), SLIDE_W, SLIDE_H, RGBColor(0x12, 0x2B, 0x57))
    frame = _textbox(slide, MARGIN, Inches(2.9), CONTENT_W, Inches(1.6))
    paragraph = frame.paragraphs[0]
    _add_runs(paragraph, title, Pt(34), WHITE, bold=True)
    if subtitle:
        paragraph = frame.add_paragraph()
        paragraph.space_before = Pt(10)
        _add_runs(paragraph, subtitle, Pt(17), RGBColor(0xC3, 0xD4, 0xF5))


def build_deck(content_path: Path, output_path: Path) -> Path:
    """Собрать PPTX из markdown-контента."""
    deck = parse_deck(content_path)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _draw_title_slide(prs, deck.meta)
    footer = deck.meta.get("footer", deck.meta.get("title", ""))

    for index, item in enumerate(deck.slides, start=2):
        slide = _new_slide(prs)
        if item.kind == "section":
            _draw_section_slide(slide, item.title, item.subtitle)
        else:
            top = _draw_header(slide, item.title, item.subtitle)
            scale = _fit_scale(item.blocks, CONTENT_BOTTOM - top, CONTENT_W)
            top += _vertical_offset(
                item.blocks, deck.base_dir, CONTENT_W, scale, CONTENT_BOTTOM - top
            )
            for block in item.blocks:
                if block.kind == "bullets":
                    top = _draw_bullets(slide, block.data, top, CONTENT_W, scale)
                elif block.kind == "table":
                    top = _draw_table(slide, block.data, top, CONTENT_W, scale)
                elif block.kind == "callout":
                    top = _draw_callout(slide, block.data, top, CONTENT_W, scale)
                elif block.kind == "code":
                    top = _draw_code(slide, block.data, top, CONTENT_W, scale)
                elif block.kind == "image":
                    src = (deck.base_dir / block.data).resolve()
                    if not src.exists():
                        raise FileNotFoundError(f"нет картинки для слайда: {src}")
                    top = _draw_image(slide, src, top, CONTENT_W, CONTENT_BOTTOM)
            if top > CONTENT_BOTTOM:
                overflow = (top - CONTENT_BOTTOM) / Inches(1)
                print(f"  ! слайд {index} «{item.title}»: переполнение на {overflow:.2f}\"")
            _draw_footer(slide, footer, index)
        if item.notes:
            slide.notes_slide.notes_text_frame.text = item.notes

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path
